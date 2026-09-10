import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel():
    wb = openpyxl.Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    # Define styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    
    section_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    section_font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
    
    title_font = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    
    pk_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fk_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    data_dict = {
        "Users & Auth": [
            {
                "table": "users",
                "desc": "Master user accounts, credentials, MFA settings, and login telemetry",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Unique user identifier"),
                    ("name", "VARCHAR(255)", "No", "—", "—", "User full name"),
                    ("email", "VARCHAR(255)", "No", "—", "UNIQUE", "User email address for authentication"),
                    ("email_verified_at", "TIMESTAMP", "Yes", "NULL", "—", "Verification timestamp"),
                    ("password", "VARCHAR(255)", "No", "—", "—", "Bcrypt hashed password"),
                    ("mfa_secret", "VARCHAR(255)", "Yes", "NULL", "—", "Encrypted Google2FA / TOTP secret"),
                    ("mfa_enabled", "BOOLEAN", "No", "false", "—", "MFA 2-factor status flag"),
                    ("is_active", "BOOLEAN", "No", "true", "—", "Account status flag (active vs disabled)"),
                    ("last_login_at", "TIMESTAMP", "Yes", "NULL", "—", "Timestamp of last successful login"),
                    ("last_login_ip", "VARCHAR(45)", "Yes", "NULL", "—", "IPv4/IPv6 address of last login"),
                    ("remember_token", "VARCHAR(100)", "Yes", "NULL", "—", "Remember me session token"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Record creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Record update timestamp")
                ]
            },
            {
                "table": "password_reset_tokens",
                "desc": "Temporary tokens for password recovery",
                "columns": [
                    ("email", "VARCHAR(255)", "No", "—", "PK", "User email address"),
                    ("token", "VARCHAR(255)", "No", "—", "—", "Hashed password reset token"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Token issuance timestamp")
                ]
            },
            {
                "table": "sessions",
                "desc": "HTTP database session storage",
                "columns": [
                    ("id", "VARCHAR(255)", "No", "—", "PK", "Unique session ID"),
                    ("user_id", "BIGINT UNSIGNED", "Yes", "NULL", "FK, INDEX", "References users.id"),
                    ("ip_address", "VARCHAR(45)", "Yes", "NULL", "—", "Client IP address"),
                    ("user_agent", "TEXT", "Yes", "NULL", "—", "Client browser user agent"),
                    ("payload", "LONGTEXT", "No", "—", "—", "Serialized session payload"),
                    ("last_activity", "INT", "No", "—", "INDEX", "Unix timestamp of last activity")
                ]
            }
        ],
        "RBAC Access Control": [
            {
                "table": "roles",
                "desc": "Spatie permission role definitions",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Unique role identifier"),
                    ("name", "VARCHAR(255)", "No", "—", "UNIQUE", "Role name (superadmin, admin, staff, customer)"),
                    ("guard_name", "VARCHAR(255)", "No", "—", "—", "Guard name (web)"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "permissions",
                "desc": "Spatie granular permission capabilities",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Unique permission identifier"),
                    ("name", "VARCHAR(255)", "No", "—", "UNIQUE", "Permission name (product.view, pos.access, etc.)"),
                    ("guard_name", "VARCHAR(255)", "No", "—", "—", "Guard name (web)"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "model_has_roles",
                "desc": "Polymorphic mapping of roles to users",
                "columns": [
                    ("role_id", "BIGINT UNSIGNED", "No", "—", "PK, FK", "References roles.id (CASCADE)"),
                    ("model_type", "VARCHAR(255)", "No", "—", "PK, INDEX", "Target model class (App\\Models\\User)"),
                    ("model_id", "BIGINT UNSIGNED", "No", "—", "PK, INDEX", "Target model ID")
                ]
            },
            {
                "table": "model_has_permissions",
                "desc": "Direct mapping of permissions to users",
                "columns": [
                    ("permission_id", "BIGINT UNSIGNED", "No", "—", "PK, FK", "References permissions.id (CASCADE)"),
                    ("model_type", "VARCHAR(255)", "No", "—", "PK, INDEX", "Target model class (App\\Models\\User)"),
                    ("model_id", "BIGINT UNSIGNED", "No", "—", "PK, INDEX", "Target model ID")
                ]
            },
            {
                "table": "role_has_permissions",
                "desc": "Mapping of permissions assigned to roles",
                "columns": [
                    ("permission_id", "BIGINT UNSIGNED", "No", "—", "PK, FK", "References permissions.id (CASCADE)"),
                    ("role_id", "BIGINT UNSIGNED", "No", "—", "PK, FK", "References roles.id (CASCADE)")
                ]
            }
        ],
        "Catalog & Inventory": [
            {
                "table": "categories",
                "desc": "Hierarchical category tree with self-referencing parent_id",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Category ID"),
                    ("name", "VARCHAR(255)", "No", "—", "—", "Category display name"),
                    ("slug", "VARCHAR(255)", "No", "—", "UNIQUE", "SEO URL slug"),
                    ("description", "TEXT", "Yes", "NULL", "—", "Category description"),
                    ("parent_id", "BIGINT UNSIGNED", "Yes", "NULL", "FK", "Self-ref FK to categories.id (SET NULL)"),
                    ("is_active", "BOOLEAN", "No", "true", "—", "Active status toggle"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "products",
                "desc": "Master product catalog with soft deletes",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Product ID"),
                    ("category_id", "BIGINT UNSIGNED", "No", "—", "FK, INDEX", "References categories.id (RESTRICT)"),
                    ("sku", "VARCHAR(255)", "No", "—", "UNIQUE", "Stock Keeping Unit"),
                    ("name", "VARCHAR(255)", "No", "—", "—", "Product title/name"),
                    ("slug", "VARCHAR(255)", "No", "—", "UNIQUE", "SEO URL slug"),
                    ("description", "TEXT", "Yes", "NULL", "—", "Full product description"),
                    ("price", "DECIMAL(10,2)", "No", "—", "—", "Base selling price"),
                    ("cost_price", "DECIMAL(10,2)", "Yes", "NULL", "—", "Cost price for margin analytics"),
                    ("is_active", "BOOLEAN", "No", "true", "INDEX", "Publishing status toggle"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp"),
                    ("deleted_at", "TIMESTAMP", "Yes", "NULL", "—", "Soft delete timestamp")
                ]
            },
            {
                "table": "product_images",
                "desc": "Product photo gallery and primary thumbnail flag",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Image ID"),
                    ("product_id", "BIGINT UNSIGNED", "No", "—", "FK", "References products.id (CASCADE)"),
                    ("path", "VARCHAR(255)", "No", "—", "—", "Storage path"),
                    ("is_primary", "BOOLEAN", "No", "false", "—", "Main thumbnail flag"),
                    ("sort_order", "INT UNSIGNED", "No", "0", "—", "Display order sequence"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "product_variations",
                "desc": "Product variations (sizes, volumes, colors, bundles)",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Variation ID"),
                    ("product_id", "BIGINT UNSIGNED", "No", "—", "FK", "References products.id (CASCADE)"),
                    ("name", "VARCHAR(255)", "No", "—", "—", "Variant label (e.g. 500ml, Size: L)"),
                    ("sku", "VARCHAR(255)", "Yes", "NULL", "—", "Variant specific SKU"),
                    ("price", "DECIMAL(10,2)", "No", "—", "—", "Variant selling price"),
                    ("cost_price", "DECIMAL(10,2)", "Yes", "NULL", "—", "Variant cost price"),
                    ("stock", "INT UNSIGNED", "No", "0", "—", "Variant stock on hand"),
                    ("is_active", "BOOLEAN", "No", "true", "—", "Active status toggle"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "inventory",
                "desc": "1:1 inventory tracking and low-stock reorder thresholds",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Inventory ID"),
                    ("product_id", "BIGINT UNSIGNED", "No", "—", "FK, UNIQUE", "References products.id (CASCADE)"),
                    ("quantity_on_hand", "INT UNSIGNED", "No", "0", "—", "Current live stock level"),
                    ("reorder_level", "INT UNSIGNED", "No", "10", "—", "Low stock alert threshold"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "stock_movements",
                "desc": "Append-only immutable stock transaction audit ledger",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Ledger ID"),
                    ("product_id", "BIGINT UNSIGNED", "No", "—", "FK, INDEX", "References products.id (CASCADE)"),
                    ("type", "ENUM", "No", "—", "—", "stock_in, stock_out, adjustment"),
                    ("quantity", "INT", "No", "—", "—", "Quantity delta (+ positive, - negative)"),
                    ("reason", "VARCHAR(255)", "Yes", "NULL", "—", "Reason / PO note / sale reference"),
                    ("performed_by", "BIGINT UNSIGNED", "No", "—", "FK", "References users.id (RESTRICT)"),
                    ("created_at", "TIMESTAMP", "No", "CURRENT_TIMESTAMP", "INDEX", "Ledger entry timestamp")
                ]
            }
        ],
        "Orders & POS": [
            {
                "table": "carts",
                "desc": "Active customer shopping cart session",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Cart ID"),
                    ("user_id", "BIGINT UNSIGNED", "No", "—", "FK, UNIQUE", "References users.id (CASCADE)"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "cart_items",
                "desc": "Items inside a user cart",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Item ID"),
                    ("cart_id", "BIGINT UNSIGNED", "No", "—", "FK, UNIQUE", "References carts.id (CASCADE)"),
                    ("product_id", "BIGINT UNSIGNED", "No", "—", "FK, UNIQUE", "References products.id (CASCADE)"),
                    ("quantity", "INT UNSIGNED", "No", "1", "—", "Quantity in cart"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "wishlists",
                "desc": "Bookmarked favorite products by customer",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Wishlist ID"),
                    ("user_id", "BIGINT UNSIGNED", "No", "—", "FK, UNIQUE", "References users.id (CASCADE)"),
                    ("product_id", "BIGINT UNSIGNED", "No", "—", "FK, UNIQUE", "References products.id (CASCADE)"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "orders",
                "desc": "Placed web & POS checkout orders",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Order ID"),
                    ("user_id", "BIGINT UNSIGNED", "No", "—", "FK", "References users.id (RESTRICT)"),
                    ("order_number", "VARCHAR(255)", "No", "—", "UNIQUE", "Human-friendly reference code"),
                    ("status", "ENUM", "No", "pending", "INDEX", "pending, processing, shipped, completed, cancelled"),
                    ("source", "VARCHAR(20)", "No", "web", "INDEX", "Origin channel (web, pos)"),
                    ("total_amount", "DECIMAL(10,2)", "No", "—", "—", "Grand total amount"),
                    ("shipping_name", "VARCHAR(255)", "No", "—", "—", "Recipient / customer name"),
                    ("shipping_phone", "VARCHAR(255)", "No", "—", "—", "Recipient contact phone"),
                    ("shipping_address", "TEXT", "No", "—", "—", "Delivery address / POS store note"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Order timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "order_items",
                "desc": "Line items of orders with historical price snapshots",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Line item ID"),
                    ("order_id", "BIGINT UNSIGNED", "No", "—", "FK", "References orders.id (CASCADE)"),
                    ("product_id", "BIGINT UNSIGNED", "No", "—", "FK", "References products.id (RESTRICT)"),
                    ("variation_id", "BIGINT UNSIGNED", "Yes", "NULL", "FK", "References product_variations.id (SET NULL)"),
                    ("variant_name", "VARCHAR(255)", "Yes", "NULL", "—", "Snapshot of variant name at purchase"),
                    ("quantity", "INT UNSIGNED", "No", "—", "—", "Quantity purchased"),
                    ("unit_price", "DECIMAL(10,2)", "No", "—", "—", "Price snapshot at purchase time"),
                    ("subtotal", "DECIMAL(10,2)", "No", "—", "—", "Line subtotal (unit_price * qty)"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            },
            {
                "table": "payments",
                "desc": "Payment transactions for orders",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Payment ID"),
                    ("order_id", "BIGINT UNSIGNED", "No", "—", "FK", "References orders.id (CASCADE)"),
                    ("method", "VARCHAR(50)", "No", "—", "—", "cash, qr_code, card, cod, online_simulation"),
                    ("status", "ENUM", "No", "pending", "—", "pending, paid, failed"),
                    ("transaction_ref", "VARCHAR(255)", "Yes", "NULL", "UNIQUE", "Gateway reference code"),
                    ("amount", "DECIMAL(10,2)", "No", "—", "—", "Amount settled"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Payment creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Payment update timestamp")
                ]
            }
        ],
        "Logs & Notifications": [
            {
                "table": "activity_logs",
                "desc": "Security and forensic action audit trail",
                "columns": [
                    ("id", "BIGINT UNSIGNED", "No", "Auto Increment", "PK", "Log ID"),
                    ("user_id", "BIGINT UNSIGNED", "Yes", "NULL", "FK, INDEX", "References users.id (SET NULL)"),
                    ("action", "VARCHAR(255)", "No", "—", "INDEX", "Action key (product.deleted, login.failed, etc.)"),
                    ("ip_address", "VARCHAR(45)", "Yes", "NULL", "—", "Client IP address"),
                    ("user_agent", "TEXT", "Yes", "NULL", "—", "Client User Agent"),
                    ("description", "JSON", "Yes", "NULL", "—", "Structured old/new payload"),
                    ("created_at", "TIMESTAMP", "No", "CURRENT_TIMESTAMP", "INDEX", "Event timestamp (append-only)")
                ]
            },
            {
                "table": "notifications",
                "desc": "System and user notifications queue",
                "columns": [
                    ("id", "CHAR(36)", "No", "—", "PK", "UUID primary key"),
                    ("type", "VARCHAR(255)", "No", "—", "—", "Notification class name"),
                    ("notifiable_type", "VARCHAR(255)", "No", "—", "INDEX", "Morph class (App\\Models\\User)"),
                    ("notifiable_id", "BIGINT UNSIGNED", "No", "—", "INDEX", "Morph ID"),
                    ("data", "TEXT", "No", "—", "—", "JSON notification payload"),
                    ("read_at", "TIMESTAMP", "Yes", "NULL", "—", "Read timestamp"),
                    ("created_at", "TIMESTAMP", "Yes", "NULL", "—", "Creation timestamp"),
                    ("updated_at", "TIMESTAMP", "Yes", "NULL", "—", "Update timestamp")
                ]
            }
        ]
    }
    
    for category_name, tables in data_dict.items():
        ws = wb.create_sheet(title=category_name)
        ws.views.sheetView[0].showGridLines = True
        
        # Title
        ws.cell(row=1, column=1, value=f"Inventory System — Data Dictionary: {category_name}").font = title_font
        ws.cell(row=2, column=1, value="Auto-generated from database migrations and Eloquent models").font = subtitle_font
        
        current_row = 4
        
        for t in tables:
            # Section Header
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=6)
            s_cell = ws.cell(row=current_row, column=1, value=f"Table: {t['table']}  —  {t['desc']}")
            s_cell.font = section_font
            s_cell.fill = section_fill
            s_cell.alignment = Alignment(vertical="center", indent=1)
            ws.row_dimensions[current_row].height = 28
            current_row += 1
            
            # Table Column Headers
            headers = ["Column Name", "Data Type", "Nullable", "Default", "Key / Index", "Description"]
            for col_num, h_text in enumerate(headers, 1):
                c = ws.cell(row=current_row, column=col_num, value=h_text)
                c.font = header_font
                c.fill = header_fill
                c.alignment = Alignment(horizontal="center" if col_num in [3, 4, 5] else "left", vertical="center")
                c.border = thin_border
            ws.row_dimensions[current_row].height = 24
            current_row += 1
            
            # Rows
            for row_data in t["columns"]:
                ws.row_dimensions[current_row].height = 20
                for col_num, val in enumerate(row_data, 1):
                    c = ws.cell(row=current_row, column=col_num, value=val)
                    c.font = Font(name="Segoe UI", size=10)
                    c.border = thin_border
                    
                    # Alignments
                    if col_num in [3, 4]:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_num == 5:
                        c.alignment = Alignment(horizontal="center", vertical="center")
                        if "PK" in val:
                            c.fill = pk_fill
                            c.font = Font(name="Segoe UI", size=10, bold=True, color="B45309")
                        elif "FK" in val:
                            c.fill = fk_fill
                            c.font = Font(name="Segoe UI", size=10, bold=True, color="0369A1")
                        elif "UNIQUE" in val:
                            c.font = Font(name="Segoe UI", size=10, bold=True, color="4338CA")
                    else:
                        c.alignment = Alignment(vertical="center")
                current_row += 1
            
            current_row += 2 # gap between tables
            
        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 22
        ws.column_dimensions['F'].width = 45

    os.makedirs("c:/laragon/www/Inventory/public/docs", exist_ok=True)
    file_path = "c:/laragon/www/Inventory/public/docs/Inventory_Data_Dictionary.xlsx"
    wb.save(file_path)
    print(f"Excel successfully created at: {file_path}")

if __name__ == "__main__":
    generate_excel()
